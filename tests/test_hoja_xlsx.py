# -*- coding: utf-8 -*-
"""La hoja de ingesta en Excel. Sin red.

El CSV no sobrevive a pasar por Excel, y el LEEME pide completar lo que falta,
que es exactamente lo que alguien hace en Excel. Cubre:

  - que el xlsx viaje en el ZIP, en los dos idiomas
  - que traiga las mismas filas que el CSV
  - que un UPC con cero adelante siga siendo texto, y lo que falta se vea
"""

import csv
import io
import zipfile

import openpyxl
import pytest

from conftest import _producto, _track
from migrador import i18n
from migrador import paquete as pq


def _productos():
    return [
        _producto(
            title="Disco",
            upc="036000291452",
            tracks=[_track("Uno", isrc="ARAAA2000001"), _track("Dos", isrc="ARAAA2000002")],
        ),
        _producto(title="Otro", upc="", tracks=[_track("Solo")]),
    ]


def _hoja(datos):
    ws = openpyxl.load_workbook(io.BytesIO(datos)).active
    assert ws is not None
    return ws


def test_trae_las_mismas_filas_que_el_csv():
    ps = _productos()
    csv_filas = list(csv.reader(io.StringIO(pq.hoja_ingesta_csv(ps, "Fulano"))))
    ws = _hoja(pq.hoja_ingesta_xlsx_bytes(ps, "Fulano"))
    xlsx_filas = [[c.value or "" for c in fila] for fila in ws.iter_rows()]
    assert xlsx_filas == csv_filas


def test_los_codigos_son_texto_y_no_pierden_ceros():
    ws = _hoja(pq.hoja_ingesta_xlsx_bytes(_productos(), "Fulano"))
    col = pq.COLUMNAS_INGESTA.index("UPC") + 1
    celda = ws.cell(row=2, column=col)
    assert celda.value == "036000291452"
    assert celda.data_type == "s"
    assert celda.number_format == "@"
    dur = ws.cell(row=2, column=pq.COLUMNAS_INGESTA.index("Duration") + 1)
    assert dur.value == "3:20" and dur.data_type == "s"


def test_lo_que_falta_va_resaltado():
    ws = _hoja(pq.hoja_ingesta_xlsx_bytes(_productos(), "Fulano"))
    faltan = [c for fila in ws.iter_rows(min_row=2) for c in fila if c.value == pq.MARCA_COMPLETAR]
    assert faltan
    assert all(c.fill.fgColor.rgb.endswith(pq.AMBAR) for c in faltan)


@pytest.mark.parametrize("idioma", ["es", "en"])
def test_viaja_en_el_zip(productos_entregable, tmp_path, idioma):
    i18n.poner_idioma(idioma)
    destino = tmp_path / "p.zip"
    pq.build_zip(productos_entregable, "Artista", str(destino), log=lambda *_: None)
    nombres = zipfile.ZipFile(destino).namelist()
    assert any(n.endswith("/" + i18n.T("paq.f_ingesta_xlsx")) for n in nombres)


def test_sin_planilla_tampoco_va(productos_entregable, tmp_path):
    destino = tmp_path / "p.zip"
    pq.build_zip(productos_entregable, "Artista", str(destino), incluir_planilla=False, log=lambda *_: None)
    assert not any(n.endswith(i18n.T("paq.f_ingesta_xlsx")) for n in zipfile.ZipFile(destino).namelist())


def test_el_leeme_pide_no_editar_el_csv_con_excel():
    for idioma in i18n.IDIOMAS:
        i18n.poner_idioma(idioma)
        leeme = pq.leeme()
        assert i18n.T("paq.f_ingesta_xlsx") in leeme
        assert "EXCEL" in leeme.upper()
