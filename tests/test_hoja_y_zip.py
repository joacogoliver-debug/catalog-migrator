# -*- coding: utf-8 -*-
"""La hoja de ingesta y las planillas nombran lo que el ZIP trae. Sin red.

En una carga masiva la distribuidora cruza cada fila con su archivo por el
nombre. La hoja decía el nombre del temporal («tidal_998877.flac») mientras el
ZIP lo guardaba como «01 - Tema.flac», y la portada figuraba como «portada.jpg»
aunque en inglés el archivo se llamara «cover.jpg», o aunque no se hubieran
pedido portadas. Cubre, en los dos idiomas:

  - que todo archivo que nombra la hoja exista en el ZIP, y al revés
  - que lo que no se incluyó quede vacío
  - que las planillas digan lo mismo
"""

import csv
import io
import zipfile

import openpyxl
import pytest

from migrador import i18n
from migrador import paquete as pq


def _armar(productos, destino, **opciones):
    pq.build_zip(productos, "Artista", str(destino), log=lambda *_: None, **opciones)
    z = zipfile.ZipFile(destino)
    raiz = z.namelist()[0].split("/")[0]
    relativos = {n[len(raiz) + 1 :] for n in z.namelist()}
    ingesta = next(n for n in z.namelist() if n.endswith(".csv"))
    filas = list(csv.DictReader(io.StringIO(z.read(ingesta).decode("utf-8-sig"))))
    return z, relativos, filas


@pytest.mark.parametrize("idioma", ["es", "en"])
def test_cada_archivo_de_la_hoja_esta_en_el_zip_y_al_reves(productos_entregable, tmp_path, idioma):
    i18n.poner_idioma(idioma)
    _z, relativos, filas = _armar(productos_entregable, tmp_path / "p.zip")

    nombrados = {f[c] for f in filas for c in ("Audio File", "Cover File") if f[c]}
    assert nombrados, "la hoja tenía que nombrar algún archivo"
    assert nombrados <= relativos, nombrados - relativos

    audios = {n for n in relativos if n.endswith((".flac", ".m4a"))}
    assert audios == {f["Audio File"] for f in filas if f["Audio File"]}


def test_la_portada_se_llama_como_en_el_zip_tambien_en_ingles(productos_entregable, tmp_path):
    i18n.poner_idioma("en")
    _z, _rel, filas = _armar(productos_entregable, tmp_path / "p.zip")
    portadas = {f["Cover File"] for f in filas if f["Cover File"]}
    assert portadas and all(p.endswith("/cover.jpg") for p in portadas)


def test_lo_que_no_se_incluyo_queda_vacio(productos_entregable, tmp_path):
    _z, relativos, filas = _armar(
        productos_entregable, tmp_path / "p.zip", incluir_audio=False, incluir_portadas=False
    )
    assert all(f["Audio File"] == "" and f["Cover File"] == "" for f in filas)
    assert not any(n.endswith((".flac", ".m4a", ".jpg")) for n in relativos)


def test_un_track_sin_audio_no_nombra_ningun_archivo(productos_entregable, tmp_path):
    _z, _rel, filas = _armar(productos_entregable, tmp_path / "p.zip")
    sin = [f for f in filas if f["Track Title"] == "Tema Sin Audio"]
    assert sin and sin[0]["Audio File"] == ""


def test_las_planillas_dicen_lo_mismo_que_el_zip(productos_entregable, tmp_path):
    z, relativos, _filas = _armar(productos_entregable, tmp_path / "p.zip")
    raiz = z.namelist()[0].split("/")[0]

    maestra = next(n for n in z.namelist() if n.endswith(".xlsx") and n.count("/") == 1)
    ws = openpyxl.load_workbook(io.BytesIO(z.read(maestra))).active
    assert ws is not None
    col = [c.value for c in ws[4]].index(i18n.T("paq.col_archivo")) + 1
    en_maestra = {str(ws.cell(row=r, column=col).value or "") for r in range(5, ws.max_row + 1)} - {""}
    assert en_maestra and en_maestra <= relativos

    # En la planilla de cada producto va el nombre solo: está en esa carpeta.
    for p in productos_entregable:
        datos = f"{raiz}/{p['folder']}/{i18n.T('paq.f_datos')}"
        ws = openpyxl.load_workbook(io.BytesIO(z.read(datos))).active
        assert ws is not None
        for r in range(5, ws.max_row + 1):
            nombre = ws.cell(row=r, column=col).value
            if nombre:
                assert f"{p['folder']}/{nombre}" in relativos
